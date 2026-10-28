"""The monthly report as an Excel workbook (openpyxl)."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from dashboards import metrics

GREEN = "1F7A55"
HEAD = PatternFill("solid", fgColor=GREEN)
TOTAL = PatternFill("solid", fgColor="E3EFE9")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Border(top=Side(style="thin", color="9DB8AA"))
MONEY = '#,##0" RWF"'
PERCENT = '0.0"%"'
NUMBER = "#,##0"
DECIMAL = "#,##0.00"

UNIT_FORMATS = {"percent": PERCENT, "money": MONEY, "count": NUMBER, "days": "0.0", "number": DECIMAL}


def _num(value):
    return float(value) if value is not None else None


def _sheet(wb, title, data, headers, rows, formats=None, widths=None, total=None):
    ws = wb.create_sheet(title)
    ws.append([f"EUCL Wellness M&E — {data['subtitle']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Generated {data['generated_at']:%d %b %Y %H:%M}"])
    ws.append([])
    ws.append(headers)
    for cell in ws[4]:
        cell.fill, cell.font = HEAD, WHITE_BOLD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    if total:
        ws.append(total)
        for cell in ws[ws.max_row]:
            cell.fill, cell.font, cell.border = TOTAL, BOLD, THIN
    for column, fmt in (formats or {}).items():
        for (cell,) in ws.iter_rows(min_row=5, min_col=column, max_col=column):
            if isinstance(fmt, str):
                cell.number_format = fmt
    for index, width in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A5"
    return ws


def build_xlsx(data, report):
    wb = Workbook()
    wb.remove(wb.active)
    k = data["kpis"]

    # ---- Summary
    ws = _sheet(wb, "Summary", data, ["Measure", "Value", f"Change from previous {data['unit']}"], [
        ["Services used (visits)", k["visits"], _num(k["visits_change"])],
        ["EUCL employees who used a service", k["participants"], _num(k["participants_change"])],
        ["Active EUCL employees", k["active_employees"], None],
        ["Participation rate (%)", _num(k["participation"]), None],
        ["Amount spent (RWF)", _num(k["spend"]), _num(k["spend_change"])],
        ["Cost per visit (RWF)", _num(k["cost_per_visit"]), None],
        [f"Still owed for this {data['unit']} (RWF)", _num(k["unpaid"]), None],
        ["Owed in total, all periods (RWF)", _num(k["owed_all"]), None],
    ], widths=[38, 18, 22])
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        label = row[0].value or ""
        row[1].number_format = MONEY if "RWF" in label else PERCENT if "%" in label else NUMBER
        row[2].number_format = '+0.0"%";-0.0"%";0"%"'
    ws.append([])
    ws.append([data["trend_title"], "Visits", "Participants", "Amount (RWF)"])
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = HEAD, WHITE_BOLD
    for t in data["trend"]:
        ws.append([t["label"], t["visits"], t["participants"], t["spend"]])
        ws.cell(ws.max_row, 4).number_format = MONEY

    # ---- Indicators: targets vs actuals (monthly reports; targets are per month)
    if data["indicators"]:
        _indicator_sheet(wb, data)

    _rest(wb, data)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _indicator_sheet(wb, data):
    ws = _sheet(wb, "Indicators", data, ["Indicator", "Unit", "Baseline", "Target", "Actual", "Achievement (%)", "Status"], [
        [
            r["indicator"].name, r["indicator"].get_unit_display(), _num(r["baseline"]), _num(r["target"]),
            _num(r["actual"]), _num(r["achievement"]), r["status_label"],
        ]
        for r in data["indicators"]
    ], formats={6: PERCENT}, widths=[44, 10, 12, 14, 14, 16, 12])
    for row_index, r in enumerate(data["indicators"], start=5):
        for column in (3, 4, 5):
            ws.cell(row_index, column).number_format = UNIT_FORMATS.get(r["indicator"].unit, DECIMAL)


def _rest(wb, data):
    # ---- Services
    services = data["services"]
    _sheet(wb, "Services", data,
           ["Service", "Company", "Covered by EUCL", "EUCL pays (RWF)", "Visits", "Employees", "Amount (RWF)", "Share (%)", f"Visits previous {data['unit']}", "Change (%)"],
           [[s["name"], s["company"], "Yes" if s["covered"] else "No", _num(s["eucl_price"]), s["visits"], s["participants"],
             _num(s["spend"]), _num(s["share"]), s["previous_visits"], _num(s["change"])] for s in services],
           formats={4: MONEY, 7: MONEY, 8: PERCENT, 10: '+0.0"%";-0.0"%";0"%"'},
           widths=[26, 24, 14, 15, 9, 11, 16, 10, 15, 11],
           total=["Total", "", "", None, sum(s["visits"] for s in services), None, _num(sum(s["spend"] for s in services)), 100 if services else None, sum(s["previous_visits"] for s in services), None])

    # ---- Companies (EUCL overall only)
    if data["all_companies"]:
        companies = data["companies"]
        _sheet(wb, "Companies", data, ["Company", "Visits", "Employees", "Amount (RWF)", "Still owed (RWF)", "Share (%)"],
               [[c["company__name"], c["visits"], c["participants"], _num(c["spend"]), _num(c["unpaid"]), _num(c["share"])] for c in companies],
               formats={4: MONEY, 5: MONEY, 6: PERCENT}, widths=[30, 10, 12, 16, 16, 10],
               total=["Total", sum(c["visits"] for c in companies), None, _num(sum(c["spend"] for c in companies)),
                      _num(sum(c["unpaid"] for c in companies)), 100 if companies else None])

    # ---- Departments
    _sheet(wb, "Departments", data, ["Department", "Active employees", "Employees who used a service", "Participation (%)", "Visits", "Amount (RWF)"],
           [[d["department"], d["employees"], d["participants"], _num(d["participation"]), d["visits"], _num(d["spend"])] for d in data["departments"]],
           formats={4: PERCENT, 6: MONEY}, widths=[24, 16, 26, 16, 10, 16])

    # ---- Employees (EUCL only)
    if data["include_names"]:
        _sheet(wb, "Most active employees", data, ["Employee", "Employee ID", "Department", "Visits", "Days", "Amount (RWF)"],
               [[f"{e['employee__first_name']} {e['employee__last_name']}".strip(), e["employee__code"], e["employee__department"],
                 e["visits"], e["days"], _num(e["spend"])] for e in data["employees"]],
               formats={6: MONEY}, widths=[28, 13, 18, 9, 8, 16])

    # ---- Every visit of the period
    visits = metrics.month_visits(data["period"], data["scope"])
    headers = ["Date", "Company", "Service", "Department", "Amount (RWF)", "Paid"]
    if data["include_names"]:
        headers[3:3] = ["Employee", "Employee ID"]
    rows = []
    for v in visits:
        row = [v.visit_date, v.company.name, v.service.name, v.employee.department, float(v.price), "Yes" if v.is_paid else "No"]
        if data["include_names"]:
            row[3:3] = [v.employee.name, v.employee.code]
        rows.append(row)
    amount_col = headers.index("Amount (RWF)") + 1
    ws = _sheet(wb, "Visits", data, headers, rows, formats={1: "dd/mm/yyyy", amount_col: MONEY},
                widths=[12, 24, 22] + ([24, 12] if data["include_names"] else []) + [18, 15, 7],
                total=["Total"] + [""] * (amount_col - 2) + [sum(r[amount_col - 1] for r in rows), ""])
