from io import BytesIO
from unittest import mock, skipUnless

from django.template.loader import render_to_string
from django.test import TestCase
from openpyxl import load_workbook

from config.testing import World, make_user

from .builders import report_data
from .models import Report

try:
    import weasyprint  # noqa: F401  (needs Pango; present in the Docker image)

    HAS_WEASYPRINT = True
except OSError:
    HAS_WEASYPRINT = False


class ReportTests(TestCase):
    def setUp(self):
        self.w = World()
        self.hr = make_user("hr", ["view_me_dashboards", "manage_employees"])
        self.boss = make_user("fit_boss", ["view_company_visits"], company=self.w.fitzone)

    def create(self, user, **data):
        self.client.force_login(user)
        return self.client.post("/reports/new/", {"month": "2026-10", "format": "xlsx", **data})

    def test_excel_report_for_all_companies(self):
        response = self.create(self.hr)
        self.assertRedirects(response, "/reports/", fetch_redirect_response=False)
        report = Report.objects.get()
        self.assertEqual(report.status, Report.DONE, report.error)  # built straight away in the tests
        self.assertTrue(report.include_names)
        download = self.client.get(f"/reports/{report.id}/download/")
        self.assertEqual(download.status_code, 200)
        self.assertIn("wellness-me-monthly-2026-10-all-companies.xlsx", download["Content-Disposition"])
        wb = load_workbook(BytesIO(b"".join(download.streaming_content)))
        self.assertEqual(wb.sheetnames, ["Summary", "Indicators", "Services", "Companies", "Departments", "Most active employees", "Visits"])
        visits = wb["Visits"]
        self.assertEqual(visits.cell(visits.max_row, 1).value, "Total")
        amount_col = [c.value for c in visits[4]].index("Amount (RWF)") + 1
        self.assertEqual(visits.cell(visits.max_row, amount_col).value, 26000)
        summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=5, max_row=12)}
        self.assertEqual(summary["Participation rate (%)"], 75.0)

    def test_company_report_has_no_names_and_no_other_companies(self):
        self.create(self.boss, company="1")
        report = Report.objects.get()
        self.assertFalse(report.include_names)
        wb = load_workbook(report.file.path)
        self.assertNotIn("Companies", wb.sheetnames)
        self.assertNotIn("Most active employees", wb.sheetnames)
        self.assertNotIn("Employee", [c.value for c in wb["Visits"][4]])
        self.assertEqual(wb["Visits"].max_row, 4 + 3 + 1)  # header row, 3 FitZone visits, total

    def test_company_staff_cannot_ask_for_other_companies(self):
        self.assertEqual(self.create(self.boss, company="2").status_code, 404)
        self.assertEqual(self.create(self.boss).status_code, 404)  # "all companies"
        self.assertFalse(Report.objects.exists())

    def test_others_cannot_download(self):
        self.create(self.hr)
        report = Report.objects.get()
        self.client.force_login(self.boss)
        self.assertEqual(self.client.get(f"/reports/{report.id}/download/").status_code, 404)

    def test_failed_report_shows_the_error(self):
        with mock.patch("reports.tasks.build_xlsx", side_effect=RuntimeError("disk full")):
            self.create(self.hr)
        report = Report.objects.get()
        self.assertEqual(report.status, Report.FAILED)
        self.assertIn("disk full", report.error)

    def test_pdf_html(self):
        """The HTML WeasyPrint prints (checked here even where WeasyPrint can't run)."""
        report = Report.objects.create(requested_by=None, period_start=self.w.month.start, format=Report.PDF, include_names=True)
        html = render_to_string("reports/report_pdf.html", {**report_data(report), "report": report})
        self.assertIn("Indicators: targets vs actuals", html)
        self.assertIn("RWF 26,000", html)
        self.assertIn("John Doe", html)

    @skipUnless(HAS_WEASYPRINT, "WeasyPrint needs Pango (installed in the Docker image)")
    def test_pdf_report(self):
        self.client.force_login(self.hr)
        self.client.post("/reports/new/", {"month": "2026-10", "format": "pdf"})
        report = Report.objects.get()
        self.assertEqual(report.status, Report.DONE, report.error)
        with report.file.open("rb") as f:
            self.assertEqual(f.read(4), b"%PDF")


class PeriodReportTests(TestCase):
    """Daily and weekly reports, and the automatic ones."""

    def setUp(self):
        self.w = World()
        self.hr = make_user("hr", ["view_me_dashboards", "manage_employees"])
        self.boss = make_user("fit_boss", ["view_company_visits"], company=self.w.fitzone)

    def create(self, user, **data):
        self.client.force_login(user)
        return self.client.post("/reports/new/", {"format": "xlsx", **data})

    def test_weekly_report(self):
        # Monday 5 to Sunday 11 October 2026: Mary (FitZone) and Paul (BlueSpa)
        self.create(self.hr, period="week", date="2026-10-07")
        report = Report.objects.get()
        self.assertEqual((report.period_type, str(report.period_start)), ("week", "2026-10-05"))
        self.assertEqual(report.status, Report.DONE, report.error)
        self.assertEqual(report.filename, "wellness-me-weekly-2026-W41-all-companies.xlsx")
        wb = load_workbook(report.file.path)
        self.assertNotIn("Indicators", wb.sheetnames)  # targets are monthly
        visits = wb["Visits"]
        amount_col = [c.value for c in visits[4]].index("Amount (RWF)") + 1
        self.assertEqual(visits.cell(visits.max_row, amount_col).value, 18000)
        summary = [row[0].value for row in wb["Summary"].iter_rows(min_row=1, max_row=30)]
        self.assertIn("Last 8 weeks", summary)

    def test_daily_report_for_a_company(self):
        self.create(self.boss, period="day", date="2026-10-02", company="1")
        report = Report.objects.get()
        self.assertEqual(report.status, Report.DONE, report.error)
        wb = load_workbook(report.file.path)
        self.assertEqual(wb["Visits"].max_row, 4 + 2 + 1)  # John's 2 visits on 2 October, then the total

    def test_no_report_for_a_period_that_has_not_started(self):
        self.create(self.hr, period="month", date="2099-01-15")
        self.assertFalse(Report.objects.exists())

    def test_the_schedule_makes_each_report_once(self):
        from .tasks import last_complete_period, scheduled_reports

        self.assertEqual(str(last_complete_period("day", today=self.w.month.start).start), "2026-09-30")
        self.assertEqual(str(last_complete_period("week", today=self.w.month.start).start), "2026-09-21")
        self.assertEqual(str(last_complete_period("month", today=self.w.month.start).start), "2026-09-01")
        scheduled_reports("week")
        scheduled_reports("week")  # again: nothing new
        reports = Report.objects.filter(is_automatic=True)
        self.assertEqual(sorted(reports.values_list("format", flat=True)), ["pdf", "xlsx"])
        self.assertTrue(all(r.include_names and r.company_id is None for r in reports))
        xlsx = reports.get(format="xlsx")
        self.assertEqual(xlsx.status, Report.DONE, xlsx.error)

    def test_automatic_reports_are_for_eucl_only(self):
        from .tasks import scheduled_reports

        scheduled_reports("day")
        self.client.force_login(self.hr)
        self.assertContains(self.client.get("/reports/"), "Automatic")
        self.client.force_login(self.boss)
        self.assertNotContains(self.client.get("/reports/"), "Automatic")
        report = Report.objects.get(format="xlsx")
        self.assertEqual(self.client.get(f"/reports/{report.id}/download/").status_code, 404)

    def test_filter_by_type(self):
        self.create(self.hr, period="day", date="2026-10-02")
        self.create(self.hr, period="week", date="2026-10-02")
        response = self.client.get("/reports/?type=week")
        self.assertEqual(len(response.context["reports"]), 1)
